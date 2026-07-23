"""Safe, bounded processing for readable attachment formats."""

from __future__ import annotations

import hashlib
import json
import math
import subprocess
import wave
import zipfile
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

MAX_EXTRACTED_CHARACTERS = 1_000_000
MAX_PDF_PAGES = 500
MAX_WORKBOOK_SHEETS = 100
MAX_WORKBOOK_CELLS = 1_000_000
MAX_PRESENTATION_SLIDES = 500
MAX_ZIP_ENTRIES = 5_000
MAX_ZIP_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
MAX_ZIP_ENTRY_BYTES = 64 * 1024 * 1024
MAX_ZIP_COMPRESSION_RATIO = 200
MAX_AUDIO_DURATION_SECONDS = 3_600
MAX_AUDIO_CHANNELS = 8
MAX_AUDIO_SAMPLE_RATE = 384_000
MAX_AUDIO_SAMPLE_WIDTH = 4
FFPROBE_TIMEOUT_SECONDS = 30
FFMPEG_TIMEOUT_SECONDS = 180

_AUDIO_FORMATS = {
    "wav": "audio/wav",
    "mp3": "audio/mpeg",
    "m4a": "audio/mp4",
    "aac": "audio/aac",
    "ogg": "audio/ogg",
    "flac": "audio/flac",
    "webm": "audio/webm",
}
_DOCUMENT_FORMATS = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
_AUDIO_SUFFIXES = {
    ".wav": "wav",
    ".wave": "wav",
    ".mp3": "mp3",
    ".m4a": "m4a",
    ".aac": "aac",
    ".ogg": "ogg",
    ".oga": "ogg",
    ".opus": "ogg",
    ".flac": "flac",
    ".webm": "webm",
}
_DOCUMENT_SUFFIXES = {f".{name}": name for name in _DOCUMENT_FORMATS}
_AUDIO_MIMES = {
    "audio/wav": "wav",
    "audio/x-wav": "wav",
    "audio/wave": "wav",
    "audio/mpeg": "mp3",
    "audio/mp3": "mp3",
    "audio/mp4": "m4a",
    "audio/x-m4a": "m4a",
    "audio/aac": "aac",
    "audio/ogg": "ogg",
    "application/ogg": "ogg",
    "audio/opus": "ogg",
    "audio/flac": "flac",
    "audio/x-flac": "flac",
    "audio/webm": "webm",
}
_DOCUMENT_MIMES = {mime: name for name, mime in _DOCUMENT_FORMATS.items()}


class AttachmentProcessingError(ValueError):
    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class AudioInfo:
    duration_ms: int
    channels: int
    sample_rate: int
    sample_width: int
    frames: int


@dataclass(frozen=True)
class ProcessingResult:
    mime_type: str
    format: str
    derived_files: dict[str, bytes] = field(default_factory=dict)
    metadata: dict[str, Any] = field(default_factory=dict)
    duration_ms: int | None = None
    channels: int | None = None
    sample_rate: int | None = None
    sample_width: int | None = None


def claimed_audio_format(mime_type: str, suffix: str) -> str | None:
    return _AUDIO_MIMES.get(mime_type) or _AUDIO_SUFFIXES.get(suffix)


def claimed_document_format(mime_type: str, suffix: str) -> str | None:
    return _DOCUMENT_MIMES.get(mime_type) or _DOCUMENT_SUFFIXES.get(suffix)


def audio_mime_type(audio_format: str) -> str:
    return _AUDIO_FORMATS[audio_format]


def document_mime_type(document_format: str) -> str:
    return _DOCUMENT_FORMATS[document_format]


def signed_audio_format(data: bytes) -> str | None:
    if len(data) >= 12 and data[:4] == b"RIFF" and data[8:12] == b"WAVE":
        return "wav"
    if data.startswith(b"fLaC"):
        return "flac"
    if data.startswith(b"OggS"):
        return "ogg"
    if data.startswith(b"\x1aE\xdf\xa3"):
        return "webm"
    if len(data) >= 2 and data[0] == 0xFF and data[1] & 0xF6 == 0xF0:
        return "aac"
    if data.startswith(b"ID3") or (
        len(data) >= 2 and data[0] == 0xFF and data[1] & 0xE0 == 0xE0
    ):
        return "mp3"
    if len(data) >= 12 and data[4:8] == b"ftyp":
        brands = data[8:32]
        if any(brand in brands for brand in (b"M4A ", b"M4B ", b"mp41", b"mp42", b"isom")):
            return "m4a"
    return None


def signed_document_format(data: bytes) -> str | None:
    if b"%PDF-" in data[:1_024]:
        return "pdf"
    if not data.startswith(b"PK\x03\x04"):
        return None
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                return None
            names = {entry.filename for entry in entries}
    except (OSError, ValueError, zipfile.BadZipFile):
        return None
    if "[Content_Types].xml" not in names:
        return None
    if "word/document.xml" in names:
        return "docx"
    if "xl/workbook.xml" in names:
        return "xlsx"
    if "ppt/presentation.xml" in names:
        return "pptx"
    return None


def validate_pcm_wav(data: bytes) -> AudioInfo:
    if len(data) < 12 or data[:4] != b"RIFF" or data[8:12] != b"WAVE":
        raise AttachmentProcessingError("attachment.invalidAudio")
    if int.from_bytes(data[4:8], "little") + 8 != len(data):
        raise AttachmentProcessingError("attachment.invalidAudio")
    try:
        with wave.open(BytesIO(data), "rb") as audio:
            channels = audio.getnchannels()
            sample_rate = audio.getframerate()
            sample_width = audio.getsampwidth()
            frames = audio.getnframes()
            if audio.getcomptype() != "NONE":
                raise AttachmentProcessingError("attachment.invalidAudio")
            frame_data = audio.readframes(frames)
            if audio.readframes(1):
                raise AttachmentProcessingError("attachment.invalidAudio")
    except AttachmentProcessingError:
        raise
    except (EOFError, OSError, ValueError, wave.Error) as exc:
        raise AttachmentProcessingError("attachment.invalidAudio") from exc
    if (
        not 1 <= channels <= MAX_AUDIO_CHANNELS
        or not 1 <= sample_rate <= MAX_AUDIO_SAMPLE_RATE
        or not 1 <= sample_width <= MAX_AUDIO_SAMPLE_WIDTH
        or frames <= 0
        or frames > sample_rate * MAX_AUDIO_DURATION_SECONDS
        or len(frame_data) != frames * channels * sample_width
    ):
        raise AttachmentProcessingError("attachment.invalidAudio")
    return AudioInfo(
        max(1, (frames * 1_000 + sample_rate - 1) // sample_rate),
        channels,
        sample_rate,
        sample_width,
        frames,
    )


def _run_media(command: list[str], data: bytes, timeout: int) -> subprocess.CompletedProcess[bytes]:
    try:
        return subprocess.run(
            command,
            input=data,
            capture_output=True,
            check=False,
            timeout=timeout,
            shell=False,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise AttachmentProcessingError("attachment.invalidAudio") from exc


def _compressed_audio(data: bytes, audio_format: str) -> ProcessingResult:
    probe = _run_media(
        [
            "ffprobe",
            "-v",
            "error",
            "-protocol_whitelist",
            "pipe",
            "-show_entries",
            "format=duration:stream=codec_type,channels,sample_rate,duration",
            "-of",
            "json",
            "pipe:0",
        ],
        data,
        FFPROBE_TIMEOUT_SECONDS,
    )
    if probe.returncode != 0 or len(probe.stdout) > 64 * 1024:
        raise AttachmentProcessingError("attachment.invalidAudio")
    try:
        payload = json.loads(probe.stdout)
        streams = [item for item in payload["streams"] if item.get("codec_type") == "audio"]
        duration_values = [payload.get("format", {}).get("duration")]
        duration_values.extend(item.get("duration") for item in streams)
        durations = [float(value) for value in duration_values if value not in {None, "N/A"}]
        channels = int(streams[0]["channels"])
        sample_rate = int(streams[0]["sample_rate"])
    except (IndexError, KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise AttachmentProcessingError("attachment.invalidAudio") from exc
    if (
        not streams
        or not 1 <= channels <= MAX_AUDIO_CHANNELS
        or not 1 <= sample_rate <= MAX_AUDIO_SAMPLE_RATE
    ):
        raise AttachmentProcessingError("attachment.invalidAudio")
    duration = max(durations) if durations else None
    if duration is not None and (
        not math.isfinite(duration) or duration <= 0 or duration > MAX_AUDIO_DURATION_SECONDS
    ):
        raise AttachmentProcessingError("attachment.invalidAudio")

    conversion = _run_media(
        [
            "ffmpeg",
            "-nostdin",
            "-v",
            "error",
            "-protocol_whitelist",
            "pipe",
            "-i",
            "pipe:0",
            "-map",
            "0:a:0",
            "-vn",
            "-t",
            str(MAX_AUDIO_DURATION_SECONDS + 0.1),
            "-ac",
            "1",
            "-ar",
            "16000",
            "-acodec",
            "pcm_s16le",
            "-f",
            "s16le",
            "pipe:1",
        ],
        data,
        FFMPEG_TIMEOUT_SECONDS,
    )
    max_pcm_bytes = MAX_AUDIO_DURATION_SECONDS * 16_000 * 2
    if (
        conversion.returncode != 0
        or not conversion.stdout
        or len(conversion.stdout) > max_pcm_bytes
        or len(conversion.stdout) % 2
    ):
        raise AttachmentProcessingError("attachment.invalidAudio")
    wav_buffer = BytesIO()
    with wave.open(wav_buffer, "wb") as wav_file:
        wav_file.setnchannels(1)
        wav_file.setsampwidth(2)
        wav_file.setframerate(16_000)
        wav_file.writeframes(conversion.stdout)
    wav_data = wav_buffer.getvalue()
    info = validate_pcm_wav(wav_data)
    if duration is not None and abs(info.duration_ms - round(duration * 1_000)) > 2_000:
        raise AttachmentProcessingError("attachment.invalidAudio")
    return ProcessingResult(
        _AUDIO_FORMATS[audio_format],
        audio_format,
        {"audio.wav": wav_data},
        duration_ms=info.duration_ms,
        channels=info.channels,
        sample_rate=info.sample_rate,
        sample_width=info.sample_width,
    )


def _append_text(parts: list[str], value: Any, count: int) -> int:
    text = str(value) if value is not None else ""
    if not text:
        return count
    count += len(text) + (1 if parts else 0)
    if count > MAX_EXTRACTED_CHARACTERS:
        raise AttachmentProcessingError("attachment.invalidDocument")
    parts.append(text)
    return count


def _zip_safety_check(data: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise AttachmentProcessingError("attachment.invalidDocument")
            names: set[str] = set()
            total = 0
            for entry in entries:
                path = PurePosixPath(entry.filename)
                if (
                    not entry.filename
                    or "\\" in entry.filename
                    or path.is_absolute()
                    or any(part in {"", ".", ".."} for part in path.parts)
                    or entry.filename in names
                    or entry.flag_bits & 0x1
                    or path.name == "vbaProject.bin"
                ):
                    raise AttachmentProcessingError("attachment.invalidDocument")
                names.add(entry.filename)
                total += entry.file_size
                if entry.file_size > MAX_ZIP_ENTRY_BYTES or total > MAX_ZIP_UNCOMPRESSED_BYTES:
                    raise AttachmentProcessingError("attachment.invalidDocument")
                if entry.file_size and entry.file_size / max(1, entry.compress_size) > MAX_ZIP_COMPRESSION_RATIO:
                    raise AttachmentProcessingError("attachment.invalidDocument")
    except AttachmentProcessingError:
        raise
    except (OSError, ValueError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise AttachmentProcessingError("attachment.invalidDocument") from exc


def _extract_pdf(data: bytes) -> tuple[str, dict[str, int]]:
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted or not 1 <= len(reader.pages) <= MAX_PDF_PAGES:
            raise AttachmentProcessingError("attachment.invalidDocument")
        parts: list[str] = []
        count = 0
        for number, page in enumerate(reader.pages, 1):
            count = _append_text(parts, f"--- Seite {number} ---", count)
            count = _append_text(parts, page.extract_text() or "", count)
    except AttachmentProcessingError:
        raise
    except Exception as exc:
        raise AttachmentProcessingError("attachment.invalidDocument") from exc
    return "\n".join(parts), {"pageCount": len(reader.pages)}


def _extract_docx(data: bytes) -> tuple[str, dict[str, int]]:
    _zip_safety_check(data)
    try:
        document = Document(BytesIO(data))
        parts: list[str] = []
        count = 0
        for paragraph in document.paragraphs:
            count = _append_text(parts, paragraph.text, count)
        for table in document.tables:
            for row in table.rows:
                count = _append_text(parts, "\t".join(cell.text for cell in row.cells), count)
    except AttachmentProcessingError:
        raise
    except Exception as exc:
        raise AttachmentProcessingError("attachment.invalidDocument") from exc
    return "\n".join(parts), {}


def _extract_xlsx(data: bytes) -> tuple[str, dict[str, int]]:
    _zip_safety_check(data)
    try:
        workbook = load_workbook(
            BytesIO(data), read_only=True, data_only=True, keep_links=False
        )
        if not 1 <= len(workbook.worksheets) <= MAX_WORKBOOK_SHEETS:
            raise AttachmentProcessingError("attachment.invalidDocument")
        parts: list[str] = []
        count = 0
        cells = 0
        for sheet in workbook.worksheets:
            count = _append_text(parts, f"--- Tabelle: {sheet.title} ---", count)
            for row in sheet.iter_rows(values_only=True):
                cells += len(row)
                if cells > MAX_WORKBOOK_CELLS:
                    raise AttachmentProcessingError("attachment.invalidDocument")
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    count = _append_text(parts, "\t".join(values), count)
        sheet_count = len(workbook.worksheets)
        workbook.close()
    except AttachmentProcessingError:
        raise
    except Exception as exc:
        raise AttachmentProcessingError("attachment.invalidDocument") from exc
    return "\n".join(parts), {"sheetCount": sheet_count}


def _extract_pptx(data: bytes) -> tuple[str, dict[str, int]]:
    _zip_safety_check(data)
    try:
        presentation = Presentation(BytesIO(data))
        if not 1 <= len(presentation.slides) <= MAX_PRESENTATION_SLIDES:
            raise AttachmentProcessingError("attachment.invalidDocument")
        parts: list[str] = []
        count = 0
        for number, slide in enumerate(presentation.slides, 1):
            count = _append_text(parts, f"--- Folie {number} ---", count)
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    count = _append_text(parts, shape.text, count)
        slide_count = len(presentation.slides)
    except AttachmentProcessingError:
        raise
    except Exception as exc:
        raise AttachmentProcessingError("attachment.invalidDocument") from exc
    return "\n".join(parts), {"slideCount": slide_count}


def process_attachment(
    data: bytes,
    *,
    kind: str,
    mime_type: str,
    filename: str,
) -> ProcessingResult | None:
    suffix = Path(filename).suffix.lower()
    if kind == "text":
        try:
            text = data.decode("utf-8", errors="strict")
        except UnicodeDecodeError as exc:
            raise AttachmentProcessingError("attachment.invalidText") from exc
        if len(text) > MAX_EXTRACTED_CHARACTERS:
            raise AttachmentProcessingError("attachment.invalidText")
        return ProcessingResult(
            mime_type if mime_type.startswith("text/") else "text/plain",
            suffix.removeprefix(".") or "text",
            {"content.txt": text.encode("utf-8")},
        )
    if kind == "audio":
        signed = signed_audio_format(data)
        claimed = claimed_audio_format(mime_type, suffix)
        audio_format = signed or claimed
        if audio_format is None or (signed is not None and claimed is not None and signed != claimed):
            raise AttachmentProcessingError("attachment.invalidAudio")
        if audio_format == "wav":
            info = validate_pcm_wav(data)
            return ProcessingResult(
                "audio/wav",
                "wav",
                duration_ms=info.duration_ms,
                channels=info.channels,
                sample_rate=info.sample_rate,
                sample_width=info.sample_width,
            )
        return _compressed_audio(data, audio_format)
    if kind == "document":
        signed = signed_document_format(data)
        claimed = claimed_document_format(mime_type, suffix)
        document_format = signed or claimed
        if document_format is None or (
            signed is not None and claimed is not None and signed != claimed
        ):
            raise AttachmentProcessingError("attachment.invalidDocument")
        extractors = {
            "pdf": _extract_pdf,
            "docx": _extract_docx,
            "xlsx": _extract_xlsx,
            "pptx": _extract_pptx,
        }
        text, metadata = extractors[document_format](data)
        encoded = text.encode("utf-8")
        return ProcessingResult(
            _DOCUMENT_FORMATS[document_format],
            document_format,
            {"content.txt": encoded},
            metadata,
        )
    return None


def derived_metadata(files: dict[str, bytes]) -> dict[str, Any]:
    if not files:
        return {}
    if len(files) != 1:
        raise ValueError("Pro Attachment wird genau eine Ableitung erwartet")
    name, data = next(iter(files.items()))
    return {
        "derivedName": name,
        "derivedSizeBytes": len(data),
        "derivedSha256": hashlib.sha256(data).hexdigest(),
    }
