"""Small valid documents for attachment processing and provenance tests."""

from io import BytesIO

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter


def document_bytes(document_format: str) -> bytes:
    if document_format == "pdf":
        return pdf_bytes()
    buffer = BytesIO()
    if document_format == "docx":
        document = Document()
        document.add_paragraph("Sicherer DOCX-Inhalt")
        document.save(buffer)
    elif document_format == "xlsx":
        workbook = Workbook()
        workbook.active.append(["Name", "Wert"])
        workbook.active.append(["MatoLoop", 42])
        workbook.save(buffer)
    elif document_format == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Sicherer PPTX-Inhalt"
        presentation.save(buffer)
    else:
        raise AssertionError("Unknown test format")
    return buffer.getvalue()


def pdf_bytes(*, pages: int = 1, password: str | None = None) -> bytes:
    buffer = BytesIO()
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=100, height=100)
    if password is not None:
        writer.encrypt(password)
    writer.write(buffer)
    return buffer.getvalue()
