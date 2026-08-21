from __future__ import annotations

import json
import stat
import subprocess
import wave
import zipfile
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter
from sqlalchemy.ext.asyncio import AsyncSession

import mato_loop_backend.attachment_processing as processing
import mato_loop_backend.attachments as attachment_module
from mato_loop_backend.attachments import (
    AttachmentError,
    AttachmentService,
    AttachmentStorage,
    classify_attachment,
)
from mato_loop_backend.models import Message
from mato_loop_backend.session_storage import SessionStorage


def wav_bytes(*, frames: int = 80, sample_rate: int = 8_000) -> bytes:
    buffer = BytesIO()
    with wave.open(buffer, "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(sample_rate)
        audio.writeframes(b"\x00\x00" * frames)
    return buffer.getvalue()


def document_bytes(document_format: str) -> bytes:
    buffer = BytesIO()
    if document_format == "pdf":
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(buffer)
    elif document_format == "docx":
        document = Document()
        document.add_paragraph("Sicherer DOCX-Inhalt")
        document.save(buffer)
    elif document_format == "xlsx":
        workbook = Workbook()
        workbook.active.append(["Name", "Wert"])
        workbook.active.append(["MatoLoop", 42])
        workbook.save(buffer)
    elif document_format == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Sicherer PPTX-Inhalt"
        presentation.save(buffer)
    else:
        raise AssertionError("Unknown test format")
    return buffer.getvalue()


def pdf_bytes(*, pages: int = 1, password: str | None = None) -> bytes:
    buffer = BytesIO()
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=100, height=100)
    if password is not None:
        writer.encrypt(password)
    writer.write(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("signature", "mime_type", "filename", "expected_mime"),
    [
        (b"ID3data", "application/octet-stream", "audio.mp3", "audio/mpeg"),
        (b"\xff\xf1data", "audio/aac", "audio.aac", "audio/aac"),
        (b"OggSdata", "audio/ogg", "audio.ogg", "audio/ogg"),
        (b"fLaCdata", "audio/flac", "audio.flac", "audio/flac"),
    ],
)
def test_common_compressed_audio_is_classified(
    signature: bytes, mime_type: str, filename: str, expected_mime: str
) -> None:
    attachment = classify_attachment(signature, mime_type, filename)

    assert attachment.kind == "audio"
    assert attachment.mime_type == expected_mime


@pytest.mark.parametrize(
    ("data", "mime_type", "filename", "media_format", "expected_mime"),
    [
        (b"\x1aE\xdf\xa3data", "audio/webm", "audio.webm", "webm", "audio/webm"),
        (b"\x00\x00\x00\x18ftypM4A data", "audio/mp4", "audio.m4a", "mp4", "audio/mp4"),
    ],
)
def test_audio_only_mp4_and_webm_are_classified_by_streams(
    monkeypatch,
    data: bytes,
    mime_type: str,
    filename: str,
    media_format: str,
    expected_mime: str,
) -> None:
    from mato_loop_backend.video_processing import MediaInfo

    monkeypatch.setattr(
        attachment_module,
        "probe_media",
        lambda _data: MediaInfo(media_format, f"video/{media_format}", 1_000, None, None, None, 0),
    )

    attachment = classify_attachment(data, mime_type, filename)

    assert (attachment.kind, attachment.mime_type) == ("audio", expected_mime)


def test_video_stream_wins_over_audio_mime_and_webm_signature(monkeypatch) -> None:
    from mato_loop_backend.video_processing import MediaInfo

    monkeypatch.setattr(
        attachment_module,
        "probe_media",
        lambda _data: MediaInfo("webm", "video/webm", 1_000, 640, 360, 0, 1),
    )

    attachment = classify_attachment(b"\x1aE\xdf\xa3data", "audio/webm", "clip.webm")

    assert (attachment.kind, attachment.mime_type) == ("video", "video/webm")


@pytest.mark.parametrize("document_format", ["pdf", "docx", "xlsx", "pptx"])
def test_modern_documents_are_classified_and_extracted(document_format: str) -> None:
    data = document_bytes(document_format)
    attachment = classify_attachment(data, "application/octet-stream", f"datei.{document_format}")

    result = processing.process_attachment(
        data,
        kind=attachment.kind,
        mime_type=attachment.mime_type,
        filename=attachment.filename or "datei",
    )

    assert attachment.kind == "document"
    assert result is not None
    assert result.format == document_format
    assert "content.txt" in result.derived_files
    if document_format in {"docx", "pptx"}:
        assert b"Sicherer" in result.derived_files["content.txt"]
    if document_format == "xlsx":
        assert b"MatoLoop" in result.derived_files["content.txt"]


def test_claimed_corrupt_readable_formats_are_not_downgraded_to_generic_files() -> None:
    claimed = classify_attachment(b"kaputt", "application/pdf", "bericht.pdf")
    assert claimed.kind == "document"
    with pytest.raises(processing.AttachmentProcessingError, match="attachment.invalidDocument"):
        processing.process_attachment(
            claimed.data,
            kind=claimed.kind,
            mime_type=claimed.mime_type,
            filename=claimed.filename or "bericht.pdf",
        )

    with pytest.raises(AttachmentError, match="attachment.invalidText"):
        classify_attachment(b"ungueltig\xff", "text/plain", "notiz.txt")


def test_text_character_limit_has_a_specific_error() -> None:
    limits = processing.AttachmentProcessingLimits(max_extracted_characters=3)

    with pytest.raises(
        processing.AttachmentProcessingError, match="attachment.textTooLarge"
    ):
        processing.process_attachment(
            b"four",
            kind="text",
            mime_type="text/csv",
            filename="logs.csv",
            limits=limits,
        )


def test_ooxml_zip_bomb_ratio_is_rejected_before_parsing() -> None:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("word/document.xml", "x" * 2_000_000)

    with pytest.raises(processing.AttachmentProcessingError, match="attachment.invalidDocument"):
        processing.process_attachment(
            buffer.getvalue(),
            kind="document",
            mime_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            filename="bombe.docx",
        )


def test_document_character_limit_is_enforced(monkeypatch) -> None:
    monkeypatch.setattr(processing, "MAX_EXTRACTED_CHARACTERS", 10)

    with pytest.raises(processing.AttachmentProcessingError, match="attachment.invalidDocument"):
        processing.process_attachment(
            document_bytes("docx"),
            kind="document",
            mime_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            filename="lang.docx",
        )


def test_pdf_page_limit_is_enforced(monkeypatch) -> None:
    monkeypatch.setattr(processing, "MAX_PDF_PAGES", 0)

    with pytest.raises(processing.AttachmentProcessingError, match="attachment.invalidDocument"):
        processing.process_attachment(
            document_bytes("pdf"),
            kind="document",
            mime_type="application/pdf",
            filename="lang.pdf",
        )


def test_empty_password_pdf_is_extracted() -> None:
    pages, metadata = processing.extract_pdf_pages(pdf_bytes(password=""))

    assert pages == [""]
    assert metadata["pageCount"] == 1


def test_password_protected_pdf_is_rejected() -> None:
    with pytest.raises(processing.AttachmentProcessingError, match="attachment.invalidDocument"):
        processing.extract_pdf_pages(pdf_bytes(password="secret"))


def test_large_pdf_manual_is_within_page_limit() -> None:
    pages, metadata = processing.extract_pdf_pages(pdf_bytes(pages=627))

    assert len(pages) == 627
    assert metadata["pageCount"] == 627


def test_compressed_audio_duration_limit_is_checked_before_conversion(monkeypatch) -> None:
    calls = 0

    def run_media(command: list[str], _data: bytes, _timeout: int):
        nonlocal calls
        calls += 1
        payload = {
            "streams": [{"codec_type": "audio", "channels": 1, "sample_rate": "8000"}],
            "format": {"duration": "3600.1"},
        }
        return subprocess.CompletedProcess(command, 0, json.dumps(payload).encode(), b"")

    monkeypatch.setattr(processing, "_run_media", run_media)

    with pytest.raises(processing.AttachmentProcessingError, match="attachment.invalidAudio"):
        processing.process_attachment(
            b"ID3too-long",
            kind="audio",
            mime_type="audio/mpeg",
            filename="lang.mp3",
        )
    assert calls == 1


async def test_compressed_audio_uses_verified_pcm_derivative_and_discards_both(
    tmp_path: Path, monkeypatch
) -> None:
    session_id = uuid4()
    sessions = SessionStorage(tmp_path / "sessions")
    sessions.provision(session_id)
    storage = AttachmentStorage(sessions)
    service = AttachmentService(storage)
    message = Message(id=uuid4(), session_id=session_id, role="user", content="Audio")
    db = AsyncMock(spec=AsyncSession)
    db.scalar.side_effect = [session_id, None]
    derived_wav = wav_bytes(frames=160, sample_rate=16_000)

    def run_media(command: list[str], _data: bytes, _timeout: int):
        if command[0] == "ffprobe":
            payload = {
                "streams": [
                    {"codec_type": "audio", "channels": 1, "sample_rate": "8000"}
                ],
                "format": {"duration": "0.01"},
            }
            return subprocess.CompletedProcess(command, 0, json.dumps(payload).encode(), b"")
        return subprocess.CompletedProcess(command, 0, b"\x00\x00" * 160, b"")

    monkeypatch.setattr(processing, "_run_media", run_media)
    original = b"ID3compressed-test-audio"

    attachment = await service.create_attachment(
        db,
        message,
        original,
        "audio/mpeg",
        kind="audio",
        filename="aufnahme.mp3",
    )

    derived_path = storage.derived_path(
        session_id, attachment.id, attachment.metadata_json["derivedPath"]
    )
    assert derived_path.read_bytes() == derived_wav
    assert stat.S_IMODE(derived_path.stat().st_mode) == 0o600
    assert await service.read_audio(attachment, session_id) == derived_wav

    derived_path.write_bytes(b"manipuliert")
    with pytest.raises(AttachmentError, match="attachment.notFound"):
        await service.read_audio(attachment, session_id)

    await service.discard(attachment, session_id)
    assert not derived_path.parent.exists()
    assert not (sessions.session_path(session_id) / attachment.storage_path).exists()


async def test_invalid_compressed_audio_is_rejected_without_persisted_files(
    tmp_path: Path, monkeypatch
) -> None:
    session_id = uuid4()
    sessions = SessionStorage(tmp_path / "sessions")
    sessions.provision(session_id)
    service = AttachmentService(AttachmentStorage(sessions))
    message = Message(id=uuid4(), session_id=session_id, role="user", content="Audio")
    db = AsyncMock(spec=AsyncSession)
    monkeypatch.setattr(
        processing,
        "_run_media",
        lambda command, data, timeout: subprocess.CompletedProcess(command, 1, b"", b"invalid"),
    )

    with pytest.raises(AttachmentError, match="attachment.invalidAudio"):
        await service.create_attachment(
            db,
            message,
            b"ID3invalid",
            "audio/mpeg",
            kind="audio",
            filename="kaputt.mp3",
        )
    assert list((sessions.session_path(session_id) / "attachments").iterdir()) == []
    assert list((sessions.session_path(session_id) / "derived").iterdir()) == []


async def test_session_quota_includes_new_derived_content(
    tmp_path: Path, monkeypatch
) -> None:
    session_id = uuid4()
    sessions = SessionStorage(tmp_path / "sessions")
    sessions.provision(session_id)
    service = AttachmentService(AttachmentStorage(sessions))
    message = Message(id=uuid4(), session_id=session_id, role="user", content="Text")
    db = AsyncMock(spec=AsyncSession)
    db.scalar.side_effect = [session_id]
    monkeypatch.setattr(attachment_module, "SESSION_ATTACHMENT_QUOTA_BYTES", 150)

    with pytest.raises(AttachmentError, match="attachment.sessionQuotaExceeded"):
        await service.create_attachment(
            db,
            message,
            b"x" * 100,
            "text/plain",
            kind="text",
            filename="lang.txt",
        )
    assert list((sessions.session_path(session_id) / "attachments").iterdir()) == []
    assert list((sessions.session_path(session_id) / "derived").iterdir()) == []
